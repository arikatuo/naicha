import json
from pathlib import Path
from openpyxl import load_workbook

folder = Path(__file__).parent
path = folder / '40品牌饮品与小料热量_核实稿.xlsx'
book = load_workbook(path, read_only=False, data_only=True)
assert book.sheetnames == ['饮品核实表', '小料核实表', '阅读说明']
drinks = json.loads((folder / 'drinks.json').read_text(encoding='utf-8'))
toppings = json.loads((folder / 'toppings.json').read_text(encoding='utf-8'))
sheet = book.worksheets[0]
assert sheet.max_row == 321 and sheet.max_column == 15
for row, original in zip(sheet.iter_rows(min_row=2, values_only=True), drinks):
    assert row[0] == original['id'] and row[1] == original['brand'] and row[2] == original['name']
    assert row[5] == original['kcal'] and row[7] == original['evidenceType']
    assert row[9] == '待核实' and row[10] is None
    assert row[11] == original['sourceUrl'] and row[12] == original['sourceNote'] and row[13] == original['verifyNote']
assert book.worksheets[1].max_row == len(toppings) + 1 == 30
for row, original in zip(book.worksheets[1].iter_rows(min_row=2, values_only=True), toppings):
    assert row[3] == original['kcal'] and row[9] == original['sourceUrl']
for s in book:
    assert len(s.tables) == 1
    for row in s.iter_rows(values_only=True):
        for value in row:
            assert value not in ['#REF!', '#VALUE!', '#DIV/0!', '#NAME?', '#N/A']
print(json.dumps({'result': 'PASS', 'brands': 40, 'drinks': 320, 'toppings': 29, 'sheets': 3, 'bytes': path.stat().st_size, 'freeze': sheet.freeze_panes}, ensure_ascii=True))
